"""Excel (XLSX) va CSV eksport: kunlik / haftalik / oylik / yillik."""
from __future__ import annotations

import csv
import io
from datetime import date, timedelta

from app.core.timeutil import local_now, period_bounds, to_utc_range, tz
from app.db.database import db
from app.domain.models import User
from app.repositories import system as sysrepo
from app.repositories import transactions as txrepo

PERIOD_LABEL = {"day": "kunlik", "week": "haftalik", "month": "oylik", "year": "yillik"}
TYPE_LABEL = {"expense": "Xarajat", "income": "Daromad"}


async def _collect(user: User, period: str, day: date | None):
    day = day or local_now(user.timezone).date()
    ps, pe = period_bounds(period, day, user.timezone)
    s, e = to_utc_range(ps, pe, user.timezone)
    bucket = {"day": "day", "week": "day", "month": "day", "year": "month"}[period]
    async with db.user_tx(user.id) as conn:
        rows = await txrepo.list_range(conn, user.id, s, e, asc=True, limit=100_000)
        totals = await txrepo.totals(conn, user.id, s, e)
        cats = await txrepo.by_category(conn, user.id, s, e, "expense")
        series = await txrepo.series(conn, user.id, user.timezone, bucket, s, e)
    async with db.system_tx() as conn:
        await sysrepo.event(conn, "data_exported", user.id, {"period": period})
    return ps, pe, rows, totals, cats, series


def _filename(period: str, ps: date, ext: str) -> str:
    return f"hisobchi_{PERIOD_LABEL[period]}_{ps.isoformat()}.{ext}"


async def export_csv(user: User, period: str, day: date | None = None) -> tuple[str, bytes]:
    ps, pe, rows, totals, _, _ = await _collect(user, period, day)
    z = tz(user.timezone)
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";")
    w.writerow(["Sana", "Vaqt", "Turi", "Kategoriya", "Izoh", "Summa (so'm)", "Manba"])
    for r in rows:
        dt = r["occurred_at"].astimezone(z)
        w.writerow([dt.strftime("%d.%m.%Y"), dt.strftime("%H:%M"), TYPE_LABEL[r["type"]], r["category_name"] or "",
                    _csv_safe(r["description"]), r["amount"], r["source"]])
    w.writerow([])
    w.writerow(["", "", "", "", "Jami xarajat", totals["expense"]])
    w.writerow(["", "", "", "", "Jami daromad", totals["income"]])
    # Excel UTF-8 ni to'g'ri ochishi uchun BOM
    return _filename(period, ps, "csv"), ("﻿" + buf.getvalue()).encode("utf-8")


def _csv_safe(v: str | None) -> str:
    """CSV/Excel formula injection himoyasi."""
    v = v or ""
    return "'" + v if v[:1] in ("=", "+", "-", "@") else v


async def export_xlsx(user: User, period: str, day: date | None = None) -> tuple[str, bytes]:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    ps, pe, rows, totals, cats, series = await _collect(user, period, day)
    z = tz(user.timezone)
    wb = Workbook()
    head_fill = PatternFill("solid", fgColor="1F6F5C")
    head_font = Font(bold=True, color="FFFFFF")

    def header(ws, cols):
        ws.append(cols)
        for c in ws[ws.max_row]:
            c.fill, c.font = head_fill, head_font
            c.alignment = Alignment(horizontal="center")

    # 1) Xulosa
    ws = wb.active
    ws.title = "Xulosa"
    ws.append([f"Hisobchi AI — {PERIOD_LABEL[period]} hisobot"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"Davr: {ps.strftime('%d.%m.%Y')} — {(pe - timedelta(days=1)).strftime('%d.%m.%Y')}"])
    ws.append([])
    ws.append(["Jami xarajat", totals["expense"]])
    ws.append(["Jami daromad", totals["income"]])
    ws.append(["Qoldiq", totals["net"]])
    ws.append(["Yozuvlar soni", totals["count"]])
    for r in range(4, 8):
        ws.cell(r, 1).font = Font(bold=True)
        ws.cell(r, 2).number_format = '#,##0 "so\'m"'
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 20

    # 2) Tranzaksiyalar
    ws = wb.create_sheet("Tranzaksiyalar")
    header(ws, ["Sana", "Vaqt", "Turi", "Kategoriya", "Izoh", "Summa (so'm)", "Manba"])
    for r in rows:
        dt = r["occurred_at"].astimezone(z)
        ws.append([dt.date(), dt.strftime("%H:%M"), TYPE_LABEL[r["type"]],
                   f"{r['category_emoji'] or ''} {r['category_name'] or ''}".strip(),
                   _csv_safe(r["description"]), r["amount"], r["source"]])
        ws.cell(ws.max_row, 1).number_format = "DD.MM.YYYY"
        ws.cell(ws.max_row, 6).number_format = "#,##0"
    last = ws.max_row
    ws.append([])
    ws.append(["", "", "", "", "Jami xarajat", totals["expense"]])
    ws.append(["", "", "", "", "Jami daromad", totals["income"]])
    for rr in (ws.max_row - 1, ws.max_row):
        ws.cell(rr, 5).font = Font(bold=True)
        ws.cell(rr, 6).font = Font(bold=True)
        ws.cell(rr, 6).number_format = "#,##0"
    for i, wdt in enumerate([12, 8, 10, 24, 32, 16, 9], 1):
        ws.column_dimensions[get_column_letter(i)].width = wdt
    ws.freeze_panes = "A2"
    if last > 1:
        ws.auto_filter.ref = f"A1:G{last}"

    # 3) Kategoriyalar + diagramma
    ws = wb.create_sheet("Kategoriyalar")
    header(ws, ["Kategoriya", "Summa (so'm)", "Ulush %", "Soni"])
    for c in cats:
        share = round(100 * c["total"] / totals["expense"], 1) if totals["expense"] else 0
        ws.append([f"{c['emoji'] or ''} {c['name'] or 'Boshqa'}".strip(), c["total"], share, c["count"]])
        ws.cell(ws.max_row, 2).number_format = "#,##0"
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 16
    if cats:
        ch = BarChart()
        ch.type = "bar"
        ch.title = "Kategoriyalar bo'yicha xarajat"
        ch.legend = None
        ch.add_data(Reference(ws, min_col=2, min_row=1, max_row=len(cats) + 1), titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=1, min_row=2, max_row=len(cats) + 1))
        ch.height, ch.width = 9, 16
        ws.add_chart(ch, "F2")

    # 4) Dinamika + line grafik
    ws = wb.create_sheet("Dinamika")
    header(ws, ["Davr", "Xarajat", "Daromad"])
    for r in series:
        ws.append([r["period"], r["expense"], r["income"]])
        ws.cell(ws.max_row, 1).number_format = "MM.YYYY" if period == "year" else "DD.MM.YYYY"
        ws.cell(ws.max_row, 2).number_format = "#,##0"
        ws.cell(ws.max_row, 3).number_format = "#,##0"
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 16
    if series:
        lc = LineChart()
        lc.title = "Xarajat dinamikasi"
        lc.add_data(Reference(ws, min_col=2, min_row=1, max_row=len(series) + 1), titles_from_data=True)
        lc.set_categories(Reference(ws, min_col=1, min_row=2, max_row=len(series) + 1))
        lc.height, lc.width = 9, 18
        ws.add_chart(lc, "E2")

    bio = io.BytesIO()
    wb.save(bio)
    return _filename(period, ps, "xlsx"), bio.getvalue()

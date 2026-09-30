"""Excel/CSV eksport testi (DB'siz — _collect almashtiriladi)."""
import asyncio
import io
import os
from datetime import date, datetime, timezone
from uuid import uuid4

os.environ.setdefault("DATA_ENCRYPTION_KEY", "a2tra2tra2tra2tra2tra2tra2tra2tra2tra2tra2s=")

from app.domain.models import User  # noqa: E402
from app.services import export  # noqa: E402

ROWS = [
    {"id": uuid4(), "type": "expense", "amount": 35000, "currency": "UZS", "category_id": 3,
     "occurred_at": datetime(2026, 9, 30, 9, 15, tzinfo=timezone.utc), "source": "voice",
     "category_name": "Taksi", "category_emoji": "🚕", "description": "taksi", "category_key": "taxi"},
    {"id": uuid4(), "type": "expense", "amount": 80000, "currency": "UZS", "category_id": 1,
     "occurred_at": datetime(2026, 9, 30, 13, 0, tzinfo=timezone.utc), "source": "text",
     "category_name": "Ovqat", "category_emoji": "🍽", "description": "=HYPERLINK(evil)", "category_key": "food"},
]


async def fake_collect(user, period, day):
    return (date(2026, 9, 1), date(2026, 10, 1), ROWS, {"expense": 115000, "income": 0, "net": -115000, "count": 2},
            [{"name": "Ovqat", "emoji": "🍽", "total": 80000, "count": 1},
             {"name": "Taksi", "emoji": "🚕", "total": 35000, "count": 1}],
            [{"period": date(2026, 9, 29), "expense": 0, "income": 0},
             {"period": date(2026, 9, 30), "expense": 115000, "income": 0}])


def _user():
    now = datetime.now(timezone.utc)
    return User(id=uuid4(), telegram_id=1, trial_ends_at=now, created_at=now, last_active_at=now)


def test_xlsx_and_csv():
    export._collect = fake_collect
    name, data = asyncio.run(export.export_xlsx(_user(), "month"))
    assert name.endswith(".xlsx") and data[:2] == b"PK"
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Xulosa", "Tranzaksiyalar", "Kategoriyalar", "Dinamika"]
    assert wb["Xulosa"]["B4"].value == 115000
    assert wb["Tranzaksiyalar"]["E3"].value.startswith("'=")    # formula injection bloklangan
    name, data = asyncio.run(export.export_csv(_user(), "month"))
    text = data.decode("utf-8-sig")
    assert "Taksi" in text and "35000" in text and "'=HYPERLINK" in text


if __name__ == "__main__":
    test_xlsx_and_csv()
    print("PASS test_xlsx_and_csv")
